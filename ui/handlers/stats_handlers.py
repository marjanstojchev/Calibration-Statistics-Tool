"""Statistics tab signal handlers  (PST-28).

Wires:
  - statsTable cell clicks  (col 0 = graph toggle, col 1 = axis toggle)
  - Group navigation: Forward / Back / Save
  - statsClearBtn, parseProtWriteStatsPushBtn, parseProtClearStatsDataBtn
  - statsGroupsTypeButtonGroup (BurnIn / Calibration radio)
"""
from __future__ import annotations

from collections import Counter
import json
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

from PySide6.QtWidgets import QTableWidgetItem, QHeaderView
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QBrush

if TYPE_CHECKING:  # pragma: no cover
    from ui.handlers.app_context import AppContext

try:
    import logconfig
except ModuleNotFoundError:
    import sys as _sys, os as _os
    _sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), '..', '..', 'shared'))
    import logconfig

_log = logconfig.mainLoadDataLogger.logger

# ---------------------------------------------------------------------------
# Module-level graph-index tracking (mirrors statsTableGraphsDict in old code)
# ---------------------------------------------------------------------------

_graphs_dict: dict = {
    'actualRemovedGraphsList': [],
    'actualCheckedGraphsDict': {-1: 'default'},
}


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _get_prot_list(parsed_dict: dict) -> list:
    return [e['parsedData'] for e in parsed_dict['protocolParsingDictList']]


def _get_protocol_paths(ctx: "AppContext") -> list[str]:
    paths: list[str] = []
    for d in ctx.pc.burnInParsedProtDict['protocolParsingDictList']:
        paths.append(d['protocolPath'])
    for d in ctx.pc.calParsedProtDict['protocolParsingDictList']:
        paths.append(d['protocolPath'])
    return paths


def _get_graph_index() -> int:
    removed = _graphs_dict['actualRemovedGraphsList']
    checked = _graphs_dict['actualCheckedGraphsDict']
    if not removed:
        return max(checked.keys()) + 1
    idx = min(removed)
    removed.remove(idx)
    return idx


def _reset_graphs_dict() -> None:
    _graphs_dict['actualRemovedGraphsList'] = []
    _graphs_dict['actualCheckedGraphsDict'] = {-1: 'default'}


def update_stats_nav_buttons(ctx: "AppContext") -> None:
    """Enable group nav buttons only after statistics have been written (table has rows)."""
    btn_id = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    grp_dict = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict
    has_groups = bool(grp_dict['limitsGroupList'])
    has_rows = ctx.qt.statsTable.rowCount() > 0
    ctx.qt.statsGroupForwardBtn.setEnabled(has_groups and has_rows)
    ctx.qt.statsGroupBackBtn.setEnabled(has_groups and has_rows)
    ctx.qt.statsGroupSaveBtn.setEnabled(has_rows)


def clear_stats_ui(ctx: "AppContext") -> None:
    """Clear the stats table, plot, group label, and selection state entirely."""
    ctx.qt.statsTable.setRowCount(0)
    ctx.mv.statsPlotFrame.clear()
    ctx.mv.statsGroupNameLabel.setText(ctx.mv.statsGroupNameString)
    _reset_graphs_dict()
    ctx.pc.burnInLimitsGroupDict['activeGroupIndex'] = -1
    ctx.pc.calLimitsGroupDict['activeGroupIndex'] = -1
    update_stats_nav_buttons(ctx)


def _collect_plot_data(parsed_list: list, limit_name: str) -> tuple[list[int], list[float]]:
    """Gather (protocol_index_list, y_value_list) for a given limit name."""
    from shared import safeparse as sp

    prot_idx_list: list[int]   = []
    plot_data:     list[float] = []

    for i, pp in enumerate(parsed_list):
        prot_idx_list.append(i)
        if (
            hasattr(pp, 'inputFilteredMeasurementLimit')
            and pp.inputFilteredMeasurementLimit is not None
            and hasattr(pp.inputFilteredMeasurementLimit, 'limitsList')
        ):
            for lim in pp.inputFilteredMeasurementLimit.limitsList:
                if lim.name == limit_name and getattr(lim, 'valueFound', 0) == 1:
                    plot_data.append(float(lim.value))
                    break
            else:
                if hasattr(pp, 'rawParsedProtocolHead') and pp.rawParsedProtocolHead:
                    for hdr, hvals in pp.rawParsedProtocolHead.items():
                        if hdr == limit_name:
                            for v in hvals:
                                if sp.checkInt(str(v)) or sp.checkDecimal(str(v)):
                                    plot_data.append(float(v))
                                    break
                            break
        else:
            if hasattr(pp, 'rawParsedProtocolHead') and pp.rawParsedProtocolHead:
                for hdr, hvals in pp.rawParsedProtocolHead.items():
                    if hdr == limit_name:
                        for v in hvals:
                            if sp.checkInt(str(v)) or sp.checkDecimal(str(v)):
                                plot_data.append(float(v))
                                break
                        break

    # align lengths — if plot_data shorter than prot_idx_list, trim index list
    if len(plot_data) < len(prot_idx_list):
        prot_idx_list = prot_idx_list[:len(plot_data)]

    return prot_idx_list, plot_data


def _plot_statistical_series(ax, values: list[float]) -> None:
    """Render values in a way that preserves exact frequencies for sparse data."""
    sorted_values = sorted(values)
    counts = Counter(sorted_values)
    unique_values = sorted(counts)

    # Sparse or low-cardinality data is better shown as exact frequency bars.
    if len(sorted_values) <= 20 or len(unique_values) <= 12:
        heights = [counts[val] for val in unique_values]

        if len(unique_values) == 1:
            center = unique_values[0]
            width = max(abs(center) * 0.05, 0.5)
        else:
            min_gap = min(
                unique_values[idx + 1] - unique_values[idx]
                for idx in range(len(unique_values) - 1)
            )
            width = max(min_gap * 0.6, 0.1)

        ax.bar(unique_values, heights, width=width, color='steelblue', edgecolor='black', alpha=0.85)

        # Show each raw value on the baseline so repeated values remain visible.
        ax.scatter(sorted_values, [0.05] * len(sorted_values), color='navy', s=20, zorder=3)
        ax.set_ylabel('Count')
        ax.set_xlabel('Measured value')
        return

    # Broader distributions still use a histogram, but with bounded bin count.
    nof_bins = max(8, min(20, int(len(sorted_values) ** 0.5) + 2))
    ax.hist(sorted_values, bins=nof_bins, color='steelblue', edgecolor='black', alpha=0.8)
    ax.scatter(sorted_values, [0.05] * len(sorted_values), color='navy', s=14, alpha=0.7, zorder=3)
    ax.set_ylabel('Frequency')
    ax.set_xlabel('Measured value')


# ---------------------------------------------------------------------------
# Graph cell colour helpers
# ---------------------------------------------------------------------------

_LIGHT_GRAPH_COLORS = {'cyan', 'yellow', 'springgreen', 'white'}


def _apply_graph_cell_color(item: QTableWidgetItem, graph_index: int) -> None:
    from widgets import getGraphColor
    color = getGraphColor(graph_index)
    item.setBackground(QBrush(QColor(color)))
    text_color = 'black' if color in _LIGHT_GRAPH_COLORS else 'white'
    item.setForeground(QBrush(QColor(text_color)))
    item.setText(str(graph_index))
    item.setData(Qt.UserRole, graph_index)


# ---------------------------------------------------------------------------
# Fill stats table
# ---------------------------------------------------------------------------

def fill_stats_table(ctx: "AppContext") -> None:
    """Populate statsTable from limitsRowListDict (after compute_limits_stats)."""
    ctx.qt.statsTable.setRowCount(0)
    col = ctx.mv.statsTableColMapping

    for row_obj in ctx.pc.limitsRowListDict['limitsRowList']:
        row_idx = ctx.qt.statsTable.rowCount()
        ctx.qt.statsTable.insertRow(row_idx)

        def _item(val: object) -> QTableWidgetItem:
            it = QTableWidgetItem(str(val) if val != 0.0 or isinstance(val, str) else '0.0')
            return it

        # Graph column — Off / not plotted by default
        graph_item = QTableWidgetItem()
        graph_item.setData(Qt.DisplayRole, 'Off')
        graph_item.setData(Qt.UserRole, -1)
        graph_item.setFlags(Qt.ItemIsEnabled)
        ctx.qt.statsTable.setItem(row_idx, col['graph'], graph_item)

        # Axis column — ax1 by default
        axis_item = QTableWidgetItem()
        axis_item.setData(Qt.DisplayRole, 'ax1')
        axis_item.setData(Qt.UserRole, 1)
        axis_item.setFlags(Qt.ItemIsEnabled)
        ctx.qt.statsTable.setItem(row_idx, col['axis'], axis_item)

        # Error
        err_str = ', '.join(str(k) for k in row_obj.error.keys()) if row_obj.error else ''
        ctx.qt.statsTable.setItem(row_idx, col['error'],      _item(err_str))
        ctx.qt.statsTable.setItem(row_idx, col['protType'],   _item(row_obj.protType))
        ctx.qt.statsTable.setItem(row_idx, col['limitName'],  _item(row_obj.limitName))
        ctx.qt.statsTable.setItem(row_idx, col['stDev'],      _item(row_obj.stDev))
        ctx.qt.statsTable.setItem(row_idx, col['mean'],       _item(row_obj.mean))
        ctx.qt.statsTable.setItem(row_idx, col['min'],        _item(row_obj.min))
        ctx.qt.statsTable.setItem(row_idx, col['max'],        _item(row_obj.max))
        ctx.qt.statsTable.setItem(row_idx, col['maxDev'],     _item(row_obj.maxDev))
        ctx.qt.statsTable.setItem(row_idx, col['upperLimit'], _item(row_obj.upperLimit))
        ctx.qt.statsTable.setItem(row_idx, col['lowerLimit'], _item(row_obj.lowerLimit))
        ctx.qt.statsTable.setItem(row_idx, col['cg'],         _item(row_obj.cg))

    # Completion list for the save dialog (default to BurnIn names)
    burn_in_names = [
        r.limitName for r in ctx.pc.limitsRowListDict['limitsRowList']
        if r.protType == 'BurnIn'
    ]
    ctx.mv.statsGroupsSaveDialog.completionList = burn_in_names
    update_stats_nav_buttons(ctx)


# ---------------------------------------------------------------------------
# Cell click handlers
# ---------------------------------------------------------------------------

def _on_stats_cell_clicked(ctx: "AppContext", row: int, col: int) -> None:
    col_map  = ctx.mv.statsTableColMapping
    protocol = _get_protocol_paths(ctx)
    if col == col_map['graph']:
        _handle_graph_click(ctx, row, protocol)
    elif col == col_map['axis']:
        _handle_axis_click(ctx, row, protocol)


def _handle_graph_click(ctx: "AppContext", row: int, protocol: list[str]) -> None:
    col_map   = ctx.mv.statsTableColMapping
    graph_item = ctx.qt.statsTable.item(row, col_map['graph'])
    axis_item  = ctx.qt.statsTable.item(row, col_map['axis'])
    name_item  = ctx.qt.statsTable.item(row, col_map['limitName'])
    type_item  = ctx.qt.statsTable.item(row, col_map['protType'])

    if graph_item is None or name_item is None:
        return

    graph_index = graph_item.data(Qt.UserRole)
    limit_name  = name_item.data(Qt.DisplayRole)
    ax          = axis_item.data(Qt.UserRole) if axis_item else 1
    rows        = ctx.pc.limitsRowListDict['limitsRowList']

    if graph_index != -1:
        # ---- Toggle OFF: remove line from plot ----------------------------
        graph_item.setBackground(QBrush(QColor('white')))
        graph_item.setForeground(QBrush(QColor('black')))
        graph_item.setText('Off')
        graph_item.setData(Qt.UserRole, -1)

        _graphs_dict['actualRemovedGraphsList'].append(graph_index)
        if graph_index in _graphs_dict['actualCheckedGraphsDict']:
            del _graphs_dict['actualCheckedGraphsDict'][graph_index]

        if row < len(rows):
            rows[row].graph = -1

        target_ax = ctx.mv.statsPlotFrame.ax1 if ax == 1 else ctx.mv.statsPlotFrame.ax2
        for i, line in enumerate(target_ax.lines):
            if line.get_label() == limit_name:
                target_ax.lines[i].remove()
                break
        ctx.mv.statsPlotFrame.ax1.legend(title='Legends', loc='upper left')
        ctx.mv.statsPlotFrame.canvas.draw()

    else:
        # ---- Toggle ON: add line to plot ----------------------------------
        graph_index = _get_graph_index()
        _graphs_dict['actualCheckedGraphsDict'][graph_index] = limit_name
        _apply_graph_cell_color(graph_item, graph_index)

        if axis_item is not None:
            axis_item.setForeground(QBrush(QColor('blue' if ax == 1 else 'red')))

        if row < len(rows):
            rows[row].graph = graph_index

        # Build plot data
        prot_type   = type_item.data(Qt.DisplayRole) if type_item else 'BurnIn'
        parsed_dict = (
            ctx.pc.calParsedProtDict
            if prot_type == 'Calibration'
            else ctx.pc.burnInParsedProtDict
        )
        parsed_list = _get_prot_list(parsed_dict)
        prot_idx_list, plot_data = _collect_plot_data(parsed_list, limit_name)

        if plot_data:
            ctx.mv.statsPlotFrame.plot(prot_idx_list, plot_data, graph_index, limit_name, protocol, ax)


def _handle_axis_click(ctx: "AppContext", row: int, protocol: list[str]) -> None:
    col_map    = ctx.mv.statsTableColMapping
    graph_item = ctx.qt.statsTable.item(row, col_map['graph'])
    name_item  = ctx.qt.statsTable.item(row, col_map['limitName'])
    axis_item  = ctx.qt.statsTable.item(row, col_map['axis'])

    if axis_item is None:
        return

    graph_index = graph_item.data(Qt.UserRole) if graph_item else -1
    limit_name  = name_item.data(Qt.DisplayRole) if name_item else ''
    ax          = axis_item.data(Qt.UserRole)
    rows        = ctx.pc.limitsRowListDict['limitsRowList']

    new_ax  = 2 if ax == 1 else 1
    axis_item.setData(Qt.UserRole, new_ax)
    axis_item.setData(Qt.DisplayRole, f'ax{new_ax}')
    axis_item.setForeground(QBrush(QColor('blue' if new_ax == 1 else 'red')))

    if row < len(rows):
        rows[row].axis = new_ax

    # If currently plotted: move line between axes
    if graph_index == -1:
        return

    src_ax = ctx.mv.statsPlotFrame.ax1 if ax == 1 else ctx.mv.statsPlotFrame.ax2
    for i, line in enumerate(src_ax.lines):
        if line.get_label() == limit_name:
            x = list(line.get_xdata())
            y = list(line.get_ydata())
            src_ax.lines[i].remove()
            prot_type   = (
                ctx.qt.statsTable.item(row, col_map['protType']).data(Qt.DisplayRole)
                if ctx.qt.statsTable.item(row, col_map['protType']) else 'BurnIn'
            )
            parsed_dict = (
                ctx.pc.calParsedProtDict
                if prot_type == 'Calibration'
                else ctx.pc.burnInParsedProtDict
            )
            prot_paths  = _get_protocol_paths(ctx)
            ctx.mv.statsPlotFrame.plot(x, y, graph_index, limit_name, prot_paths, new_ax)
            break


# ---------------------------------------------------------------------------
# Clear selected stats data (plot + table selection state)
# ---------------------------------------------------------------------------

def clear_sel_stats_data(ctx: "AppContext") -> None:
    """Clear the plot and reset all graph/axis selections in statsTable."""
    ctx.pc.burnInLimitsGroupDict['activeGroupIndex'] = -1
    ctx.pc.calLimitsGroupDict['activeGroupIndex'] = -1
    ctx.mv.statsPlotFrame.clear()
    col_map = ctx.mv.statsTableColMapping
    rows    = ctx.pc.limitsRowListDict['limitsRowList']

    for row_idx in range(ctx.qt.statsTable.rowCount()):
        graph_item = ctx.qt.statsTable.item(row_idx, col_map['graph'])
        axis_item  = ctx.qt.statsTable.item(row_idx, col_map['axis'])

        if graph_item is not None and graph_item.data(Qt.UserRole) != -1:
            graph_item.setData(Qt.UserRole, -1)
            graph_item.setData(Qt.DisplayRole, 'Off')
            graph_item.setBackground(QBrush(QColor('white')))
            graph_item.setForeground(QBrush(QColor('black')))
            if row_idx < len(rows):
                rows[row_idx].graph = -1

        if axis_item is not None:
            axis_item.setForeground(QBrush(QColor('black')))
            axis_item.setData(Qt.UserRole, 1)
            axis_item.setData(Qt.DisplayRole, 'ax1')
            if row_idx < len(rows):
                rows[row_idx].axis = 1

    ctx.mv.statsGroupNameLabel.setText(ctx.mv.statsGroupNameString)
    _reset_graphs_dict()


# ---------------------------------------------------------------------------
# Group activation
# ---------------------------------------------------------------------------

def set_limits_group_active(ctx: "AppContext", limits_group) -> None:
    """Activate a LimitsGroup: highlight its rows and plot them."""
    protocol = _get_protocol_paths(ctx)
    clear_sel_stats_data(ctx)

    col_map = ctx.mv.statsTableColMapping
    rows    = ctx.pc.limitsRowListDict['limitsRowList']

    prot_type_label = limits_group.protType
    ctx.mv.statsGroupNameLabel.setText(f'{prot_type_label}: {limits_group.name}')

    parsed_dict = (
        ctx.pc.calParsedProtDict
        if prot_type_label == 'Calibration'
        else ctx.pc.burnInParsedProtDict
    )
    parsed_list = _get_prot_list(parsed_dict)

    for match_dict in limits_group.limitsRowMatchDictList:
        limit_row = match_dict['row']
        row_idx   = limit_row.index

        if row_idx >= ctx.qt.statsTable.rowCount():
            continue

        axis_item  = ctx.qt.statsTable.item(row_idx, col_map['axis'])
        ax         = axis_item.data(Qt.UserRole) if axis_item else 1

        graph_index = _get_graph_index()
        _graphs_dict['actualCheckedGraphsDict'][graph_index] = limit_row.limitName

        graph_item = ctx.qt.statsTable.item(row_idx, col_map['graph'])
        if graph_item is not None:
            _apply_graph_cell_color(graph_item, graph_index)

        if axis_item is not None:
            axis_item.setForeground(QBrush(QColor('blue' if ax == 1 else 'red')))

        if row_idx < len(rows):
            rows[row_idx].graph = graph_index

        prot_idx_list, plot_data = _collect_plot_data(parsed_list, limit_row.limitName)
        if plot_data:
            ctx.mv.statsPlotFrame.plot(
                prot_idx_list, plot_data, graph_index, limit_row.limitName, protocol, ax
            )


# ---------------------------------------------------------------------------
# Group navigation
# ---------------------------------------------------------------------------

def _on_show_groups_type(ctx: "AppContext") -> None:
    btn_id   = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    grp_dict = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict
    active   = grp_dict['activeGroupIndex']
    if active == -1:
        clear_sel_stats_data(ctx)
    else:
        set_limits_group_active(ctx, grp_dict['limitsGroupList'][active])
    update_stats_nav_buttons(ctx)


def _on_group_forward(ctx: "AppContext") -> None:
    btn_id   = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    grp_dict = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict
    lst      = grp_dict['limitsGroupList']
    if not lst:
        return
    idx = grp_dict['activeGroupIndex']
    idx = 0 if idx >= len(lst) - 1 else idx + 1
    set_limits_group_active(ctx, lst[idx])
    grp_dict['activeGroupIndex'] = idx


def _on_group_back(ctx: "AppContext") -> None:
    btn_id   = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    grp_dict = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict
    lst      = grp_dict['limitsGroupList']
    if not lst:
        return
    idx = grp_dict['activeGroupIndex']
    idx = len(lst) - 1 if idx <= 0 else idx - 1
    set_limits_group_active(ctx, lst[idx])
    grp_dict['activeGroupIndex'] = idx


# ---------------------------------------------------------------------------
# Save group dialog
# ---------------------------------------------------------------------------

def _on_group_save(ctx: "AppContext") -> None:
    """Open LimitsGroupsSaveDialog pre-populated with currently visible rows."""
    dialog    = ctx.mv.statsGroupsSaveDialog
    btn_id    = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    prot_type = 'BurnIn' if btn_id == 0 else 'Calibration'
    grp_dict  = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict

    # Reset dialog
    dialog.errorLabel.setText('')
    dialog.groupNameLEdit.clear()
    dialog.saveToFileCheckBox.setChecked(False)
    dialog.table.setRowCount(0)
    dialog.comboList = []
    dialog.initTable()

    # Completion list
    completion = [
        r.limitName for r in ctx.pc.limitsRowListDict['limitsRowList']
        if r.protType == prot_type
    ]
    dialog.completionList = completion

    # Pre-fill with rows that are currently graphed for this prot_type
    col_map    = ctx.mv.statsTableColMapping
    table_row  = 0
    active_idx = grp_dict['activeGroupIndex']
    active_matches: list[dict] = []
    if active_idx != -1:
        active_matches = grp_dict['limitsGroupList'][active_idx].limitsRowMatchDictList

    rows = ctx.pc.limitsRowListDict['limitsRowList']
    for r_idx in range(ctx.qt.statsTable.rowCount()):
        if r_idx >= len(rows):
            break
        row_obj = rows[r_idx]
        if row_obj.graph == -1 or row_obj.protType != prot_type:
            continue

        while table_row >= dialog.table.rowCount():
            dialog.addLine()

        limit_item = dialog.table.item(table_row, 0)
        if limit_item is not None:
            limit_item.setText(row_obj.limitName)
            comp_dict = {'completion_ls': completion, 'row': row_obj}
            limit_item.setData(Qt.UserRole, comp_dict)

        combo = dialog.table.cellWidget(table_row, 1)
        if combo:
            match_val = 'ExactMatch'
            for m in active_matches:
                if m['row'].limitName == row_obj.limitName:
                    match_val = m.get('match', 'ExactMatch')
                    break
            idx = 0 if match_val == 'ExactMatch' else 1
            combo.setCurrentIndex(idx)

        table_row += 1

    # Wire Save button once per dialog open to avoid duplicate saves.
    prev_cb = getattr(dialog, '_save_callback', None)
    if prev_cb is not None:
        try:
            dialog.buttonBox.accepted.disconnect(prev_cb)
        except Exception:
            pass

    callback = lambda: _dialog_save(ctx)
    dialog._save_callback = callback
    dialog.buttonBox.accepted.connect(callback)
    dialog.exec_()


def _dialog_save(ctx: "AppContext") -> None:
    """Validate and save the new LimitsGroup from the save dialog."""
    from shared.parseProtocols import LimitsGroup

    dialog    = ctx.mv.statsGroupsSaveDialog
    btn_id    = ctx.qt.statsGroupsTypeButtonGroup.checkedId()
    prot_type = 'BurnIn' if btn_id == 0 else 'Calibration'
    grp_dict  = ctx.pc.burnInLimitsGroupDict if btn_id == 0 else ctx.pc.calLimitsGroupDict

    group_name = dialog.groupNameLEdit.text().strip()
    if not group_name:
        dialog.errorLabel.setText('Empty Group Name field!')
        return

    new_group           = LimitsGroup()
    new_group.name      = group_name
    new_group.protType  = prot_type

    rows = ctx.pc.limitsRowListDict['limitsRowList']
    for r in range(dialog.table.rowCount()):
        limit_item = dialog.table.item(r, 0)
        if limit_item is None:
            continue
        lim_name = limit_item.text().strip()
        if not lim_name:
            continue
        combo = dialog.table.cellWidget(r, 1)
        match = combo.currentText() if combo else 'ExactMatch'

        if match == 'Contains':
            for row_obj in rows:
                if lim_name in row_obj.limitName and row_obj.protType == prot_type:
                    new_group.limitsRowMatchDictList.append({'row': row_obj, 'match': match})
        else:
            for row_obj in rows:
                if row_obj.limitName == lim_name and row_obj.protType == prot_type:
                    new_group.limitsRowMatchDictList.append({'row': row_obj, 'match': match})
                    break

    if not new_group.limitsRowMatchDictList:
        dialog.errorLabel.setText('No valid limits selected for this group.')
        return

    grp_dict['limitsGroupList'].append(new_group)
    grp_dict['activeGroupIndex'] = len(grp_dict['limitsGroupList']) - 1

    if dialog.saveToFileCheckBox.isChecked():
        try:
            _save_group_to_yaml(ctx, new_group)
        except Exception as exc:
            dialog.errorLabel.setText(f'Failed to save YAML: {exc}')
            return

    set_limits_group_active(ctx, new_group)
    update_stats_nav_buttons(ctx)
    dialog.accept()
    _log.info(f"Saved new limits group '{group_name}' ({prot_type})")


def _save_group_to_yaml(ctx: "AppContext", limits_group) -> None:
    """Persist a newly created limits group to the active YAML file path."""
    from shared import auxiliary as aux
    from shared.parseProtocols import ParseInputLimitsGroups

    prot_type = limits_group.protType
    parser = ParseInputLimitsGroups()

    if prot_type == 'BurnIn':
        file_path = ctx.qt.parseProtProtDataLimGroupsBurnInLabel.text()
        if not file_path or file_path == 'No file set.':
            file_path = aux.burnInLimGroupsFilePath
            ctx.mv.set_burn_in_lim_groups_file(file_path)

        parser.loadBurnInInputLimitsGroupsDictYAML(file_path)
        if parser.burnInErrorStatus == 1:
            raise ValueError(f'Failed to load BurnIn limits groups YAML: {file_path}')

        parser.saveToActualBurnInLimitsGroupsDict(limits_group, alreadyExtended=0)
        parser.saveLimitsGroupsDictYAML(parser.burnInInputLimitsGroupsDict, file_path)
    else:
        file_path = ctx.qt.parseProtProtDataLimGroupsCalLabel.text()
        if not file_path or file_path == 'No file set.':
            file_path = aux.calLimGroupsFilePath
            ctx.mv.set_cal_lim_groups_file(file_path)

        parser.loadCalInputLimitsGroupsDictYAML(file_path)
        if parser.calErrorStatus == 1:
            raise ValueError(f'Failed to load Calibration limits groups YAML: {file_path}')

        parser.saveToActualCalLimitsGroupsDict(limits_group, alreadyExtended=0)
        parser.saveLimitsGroupsDictYAML(parser.calInputLimitsGroupsDict, file_path)

    _log.info(f"Persisted limits group '{limits_group.name}' to: {file_path}")


def _on_stats_input_config_load(ctx: "AppContext") -> None:
    """Load and validate the statistics input config JSON file path."""
    try:
        if ctx.qt.statsInputConfFDialog.exec_() != 1:
            return

        selected = ctx.qt.statsInputConfFDialog.selectedFiles()
        if not selected:
            return

        file_path = selected[0]
        with open(file_path, 'r', encoding='utf-8') as handle:
            loaded = json.load(handle)

        if not isinstance(loaded, dict):
            raise ValueError('Invalid statistics config: root must be a JSON object.')

        ctx.mv.set_stats_input_conf_path(file_path)
        _log.info(f'Statistics input config loaded: {file_path}')
    except Exception as exc:
        _log.error(f'Failed to load statistics config: {exc}')


def _on_stats_group_pptx_export(ctx: "AppContext") -> None:
    """Export the current statistics plot to a one-slide PPTX in output/."""
    try:
        from pptx import Presentation
        from pptx.util import Inches

        output_dir = Path('output')
        output_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        png_path = output_dir / f'stats_plot_{stamp}.png'
        pptx_path = output_dir / f'stats_plot_{stamp}.pptx'

        # Persist current matplotlib canvas and embed it in a single-slide deck.
        ctx.mv.statsPlotFrame.figure.savefig(str(png_path), dpi=180, bbox_inches='tight')

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = 'Statistics Plot Export'
        slide.shapes.add_picture(str(png_path), Inches(0.5), Inches(1.3), width=Inches(9.0))
        prs.save(str(pptx_path))

        _log.info(f'Statistics PPTX exported: {pptx_path}')
    except Exception as exc:
        _log.error(f'Failed to export statistics PPTX: {exc}')


def _on_statistical_series(ctx: "AppContext") -> None:
    """Render a histogram for the selected (or first active) limits row."""
    col_map = ctx.mv.statsTableColMapping
    selected = ctx.qt.statsTable.selectedIndexes()
    row_idx = selected[0].row() if selected else -1

    if row_idx < 0:
        for i, row_obj in enumerate(ctx.pc.limitsRowListDict['limitsRowList']):
            if row_obj.graph != -1:
                row_idx = i
                break

    if row_idx < 0 or row_idx >= ctx.qt.statsTable.rowCount():
        _log.warning('Statistical Series: select or enable at least one limit row first.')
        return

    name_item = ctx.qt.statsTable.item(row_idx, col_map['limitName'])
    type_item = ctx.qt.statsTable.item(row_idx, col_map['protType'])
    if name_item is None or type_item is None:
        _log.warning('Statistical Series: invalid row selection.')
        return

    limit_name = str(name_item.data(Qt.DisplayRole))
    prot_type = str(type_item.data(Qt.DisplayRole))
    parsed_dict = ctx.pc.calParsedProtDict if prot_type == 'Calibration' else ctx.pc.burnInParsedProtDict
    parsed_list = _get_prot_list(parsed_dict)
    _, values = _collect_plot_data(parsed_list, limit_name)
    row_obj = ctx.pc.limitsRowListDict['limitsRowList'][row_idx]

    if not values:
        _log.warning(f'Statistical Series: no values available for {limit_name}.')
        return

    ctx.mv.statsPlotFrame.clear()
    ax = ctx.mv.statsPlotFrame.ax1
    _plot_statistical_series(ax, values)

    ax.set_title(f'Statistical Series - {limit_name} (n={len(values)})')
    ax.axvline(row_obj.mean, color='darkorange', linestyle='--', linewidth=1.5, label=f'Mean {row_obj.mean:.3f}')
    if row_obj.lowerLimit != 0.0 or row_obj.upperLimit != 0.0:
        ax.axvline(row_obj.lowerLimit, color='crimson', linestyle=':', linewidth=1.2, label=f'Lower {row_obj.lowerLimit:.3f}')
        ax.axvline(row_obj.upperLimit, color='forestgreen', linestyle=':', linewidth=1.2, label=f'Upper {row_obj.upperLimit:.3f}')

    ax.legend(loc='best')
    ctx.mv.statsPlotFrame.canvas.draw()
    _log.info(f'Statistical Series rendered for: {limit_name}')


# ---------------------------------------------------------------------------
# Full clear (stats table + plot + state dicts)
# ---------------------------------------------------------------------------

def _on_stats_clear(ctx: "AppContext") -> None:
    """Clear statsTable and the plot entirely."""
    clear_stats_ui(ctx)


# ---------------------------------------------------------------------------
# Write statistic (compute + fill tables)
# ---------------------------------------------------------------------------

def write_statistic(ctx: "AppContext") -> None:
    """Compute stats and populate the statistics tables.

    Called by the 'Parse / Write Statistic' button.
    """
    ctx.pc.compute_limits_stats()
    fill_stats_table(ctx)
    _fill_stats_protocol_tables(ctx)

    # Select the Statistics tab directly instead of relying on a fixed tab index.
    ctx.qt.mainTabWidget.setCurrentWidget(ctx.qt.statsTab)
    _log.info('write_statistic: complete.')


def _fill_stats_protocol_tables(ctx: "AppContext") -> None:
    """Populate parseProtStatsParsedTable and parseProtStatsFailedTable."""
    for tbl, dicts, labels in (
        (
            ctx.qt.parseProtStatsParsedTable,
            [ctx.pc.burnInParsedProtDict, ctx.pc.calParsedProtDict],
            ['BurnIn', 'Calibration'],
        ),
        (
            ctx.qt.parseProtStatsFailedTable,
            [ctx.pc.burnInFailedProtDict, ctx.pc.calFailedProtDict],
            ['BurnIn', 'Calibration'],
        ),
    ):
        tbl.setRowCount(0)
        hdr = tbl.horizontalHeader()
        hdr.setSectionResizeMode(0, QHeaderView.Interactive)
        tbl.setColumnWidth(0, 800)
        hdr.setSectionResizeMode(1, QHeaderView.Interactive)
        tbl.setColumnWidth(1, 110)
        hdr.setSectionResizeMode(2, QHeaderView.Stretch)

        for parsed_dict, p_label in zip(dicts, labels):
            for entry in parsed_dict['protocolParsingDictList']:
                r = tbl.rowCount()
                tbl.insertRow(r)
                tbl.setItem(r, 0, QTableWidgetItem(entry.get('protocolPath', '')))
                tbl.setItem(r, 1, QTableWidgetItem(entry.get('type', '')))
                tbl.setItem(r, 2, QTableWidgetItem(p_label))


# ---------------------------------------------------------------------------
# init_widgets / wire
# ---------------------------------------------------------------------------

def init_widgets(ctx: "AppContext") -> None:
    ctx.qt.statsGroupsTypeBurnInRadio.setChecked(True)
    # Let cursor click on statsPlotFrame log the file name
    ctx.mv.statsPlotFrame._on_cursor_select = lambda name: _log.info(f'Plot cursor: {name}')
    update_stats_nav_buttons(ctx)


def wire(ctx: "AppContext") -> None:
    """Connect all statistics-tab Qt signals to their handlers."""
    # Parse tab buttons
    ctx.qt.parseProtWriteStatsPushBtn.clicked.connect(
        lambda: write_statistic(ctx)
    )
    ctx.qt.parseProtClearStatsDataBtn.clicked.connect(
        lambda: _on_stats_clear(ctx)
    )

    # Statistics tab buttons
    ctx.qt.statsClearBtn.clicked.connect(
        lambda: clear_sel_stats_data(ctx)
    )
    ctx.qt.statsGroupForwardBtn.clicked.connect(
        lambda: _on_group_forward(ctx)
    )
    ctx.qt.statsGroupBackBtn.clicked.connect(
        lambda: _on_group_back(ctx)
    )
    ctx.qt.statsGroupSaveBtn.clicked.connect(
        lambda: _on_group_save(ctx)
    )
    ctx.qt.statsSeriesDistBtn.clicked.connect(
        lambda: _on_statistical_series(ctx)
    )
    ctx.qt.statsInputConfBtn.clicked.connect(
        lambda: _on_stats_input_config_load(ctx)
    )
    ctx.qt.statsGroupPptxBtn.clicked.connect(
        lambda: _on_stats_group_pptx_export(ctx)
    )

    # Group type radio buttons
    ctx.qt.statsGroupsTypeButtonGroup.buttonClicked.connect(
        lambda _btn: _on_show_groups_type(ctx)
    )

    # Cell click
    ctx.qt.statsTable.cellClicked.connect(
        lambda r, c: _on_stats_cell_clicked(ctx, r, c)
    )
