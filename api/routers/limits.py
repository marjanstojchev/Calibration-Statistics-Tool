from fastapi import HTTPException, APIRouter
from pydantic import BaseModel
import os
from typing import Optional

router = APIRouter()

# Module-level storage so GET /rows and GET /groups can read after /compute
_last_controller = None
_last_limits_type: str = "burn_in"


class ComputeStatsRequest(BaseModel):
    """Combined request for limits config + parsing options."""
    limits_type: str = "burn_in"          # "burn_in" or "calibration"
    input_limits_file_path: str            # Server-side path to limits YAML
    extend_input_limits: bool = False
    limits_groups_file_path: str           # Server-side path to groups YAML
    extend_limits_groups: bool = False
    active_limits_group: str = ""


def _row_to_dict(row) -> dict:
    return {
        "index":       row.index,
        "limitName":   row.limitName,
        "protType":    row.protType,
        "lowerLimit":  row.lowerLimit,
        "upperLimit":  row.upperLimit,
        "mean":        row.mean,
        "min":         row.min,
        "max":         row.max,
        "stDev":       row.stDev,
        "maxDev":      row.maxDev,
        "cg":          row.cg,
        "error":       row.error,
    }


def _group_to_dict(group) -> dict:
    return {
        "name":     group.name,
        "protType": group.protType,
        "rows":     group.limitsRowMatchDictList,
    }


@router.post("/compute")
def compute_stats(request: ComputeStatsRequest):
    """Parse limits files and compute full statistics.

    Flow: parse_protocol_limits → build_limits_groups → build_limits_rows →
    compute_limits_stats.  Results are stored in memory and readable via
    GET /rows and GET /groups.

    Requires /api/protocols/parse to have been called first.
    """
    global _last_controller, _last_limits_type

    from shared.parse_controller import ParseController
    import shared.load_protocols as lp

    parsed_data = getattr(lp, 'sysDiagParsedDataList', [])
    if not parsed_data:
        raise HTTPException(status_code=422, detail="No parsed data. Run /api/protocols/parse first.")

    for path in [request.input_limits_file_path, request.limits_groups_file_path]:
        if not os.path.isfile(path):
            raise HTTPException(status_code=400, detail=f"File not found: {path}")

    is_burn_in = request.limits_type == "burn_in"

    controller = ParseController()
    controller.parse_protocol_limits(
        burn_in_limits_path=request.input_limits_file_path,
        cal_limits_path=request.input_limits_file_path,
        include_burn_in=is_burn_in,
        include_cal=not is_burn_in,
        extend_burn_in=request.extend_input_limits,
        extend_cal=request.extend_input_limits,
    )
    controller.build_limits_groups(
        burn_in_lim_groups_path=request.limits_groups_file_path,
        cal_lim_groups_path=request.limits_groups_file_path,
        include_burn_in=is_burn_in,
        include_cal=not is_burn_in,
        extend_burn_in_groups=request.extend_limits_groups,
        extend_cal_groups=request.extend_limits_groups,
    )
    controller.build_limits_rows()
    controller.compute_limits_stats()

    _last_controller = controller
    _last_limits_type = request.limits_type

    return {
        "status": "ok",
        "limits_type": request.limits_type,
        "row_count": len(controller.limitsRowListDict["limitsRowList"]),
    }


@router.get("/rows")
def get_rows():
    """Return all computed limit rows with statistics.

    Must call POST /compute first.
    """
    if _last_controller is None:
        raise HTTPException(status_code=422, detail="No limits computed. Call POST /compute first.")

    rows = _last_controller.limitsRowListDict["limitsRowList"]
    return {
        "limits_type": _last_limits_type,
        "row_count": len(rows),
        "rows": [_row_to_dict(r) for r in rows],
    }


@router.get("/groups")
def get_groups():
    """Return all limits groups (with their member rows).

    Must call POST /compute first.
    """
    if _last_controller is None:
        raise HTTPException(status_code=422, detail="No limits computed. Call POST /compute first.")

    if _last_limits_type == "burn_in":
        groups = _last_controller.burnInLimitsGroupDict.get("limitsGroupList", [])
    else:
        groups = _last_controller.calLimitsGroupDict.get("limitsGroupList", [])

    return {
        "limits_type": _last_limits_type,
        "group_count": len(groups),
        "groups": [_group_to_dict(g) for g in groups],
    }


@router.delete("/clear")
def clear_limits():
    """Clear all computed limits statistics from memory.

    Equivalent to the 'Clear Statistics Data' button in the Statistics tab.
    This resets the controller state but does NOT clear parsed protocols.
    """
    global _last_controller, _last_limits_type
    _last_controller = None
    _last_limits_type = "burn_in"
    return {"status": "cleared"}


# ---------------------------------------------------------------------------
# Additional statistics features
# ---------------------------------------------------------------------------


class SaveGroupRequest(BaseModel):
    """Request to create and save a new limits group."""
    group_name: str
    limits_type: str = "burn_in"  # "burn_in" or "calibration"
    limit_names: list[str]        # List of limit names to include
    match_type: str = "ExactMatch"  # "ExactMatch" or "Contains"
    save_to_yaml: bool = True


@router.post("/groups/save")
def save_limits_group(request: SaveGroupRequest):
    """Create a new limits group and optionally persist to YAML file.

    Equivalent to the 'Save Group' dialog in the Statistics tab.
    """
    if _last_controller is None:
        raise HTTPException(status_code=422, detail="No limits computed. Call POST /compute first.")

    # Lazy-load to avoid Qt initialization at module import
    from shared.parseProtocols import LimitsGroup, ParseInputLimitsGroups
    import shared.auxiliary as aux

    # Create new group
    new_group = LimitsGroup()
    new_group.name = request.group_name
    new_group.protType = "BurnIn" if request.limits_type == "burn_in" else "Calibration"

    # Find matching rows
    rows = _last_controller.limitsRowListDict.get("limitsRowList", [])
    for limit_name in request.limit_names:
        for row_obj in rows:
            if row_obj.protType != new_group.protType:
                continue
            
            if request.match_type == "Contains":
                if limit_name in row_obj.limitName:
                    new_group.limitsRowMatchDictList.append({"row": row_obj, "match": request.match_type})
            else:  # ExactMatch
                if row_obj.limitName == limit_name:
                    new_group.limitsRowMatchDictList.append({"row": row_obj, "match": request.match_type})
                    break

    if not new_group.limitsRowMatchDictList:
        raise HTTPException(status_code=400, detail="No valid limits found matching the provided names.")

    # Add to active groups dict
    grp_dict = (
        _last_controller.burnInLimitsGroupDict
        if request.limits_type == "burn_in"
        else _last_controller.calLimitsGroupDict
    )
    grp_dict["limitsGroupList"].append(new_group)
    grp_dict["activeGroupIndex"] = len(grp_dict["limitsGroupList"]) - 1

    # Optionally save to YAML
    yaml_path = None
    if request.save_to_yaml:
        parser = ParseInputLimitsGroups()
        
        if request.limits_type == "burn_in":
            yaml_path = aux.burnInLimGroupsFilePath
            parser.loadBurnInInputLimitsGroupsDictYAML(yaml_path)
            if parser.burnInErrorStatus == 1:
                raise HTTPException(
                    status_code=500,
                    detail=f"Failed to load BurnIn limits groups YAML: {yaml_path}"
                )
            parser.saveToActualBurnInLimitsGroupsDict(new_group, alreadyExtended=0)
            parser.saveLimitsGroupsDictYAML(parser.burnInInputLimitsGroupsDict, yaml_path)
        else:
            yaml_path = aux.calLimGroupsFilePath
            parser.loadCalInputLimitsGroupsDictYAML(yaml_path)
            if parser.calErrorStatus == 1:
                raise HTTPException(
                    status_code=500,
                    detail=f"Failed to load Calibration limits groups YAML: {yaml_path}"
                )
            parser.saveToActualCalLimitsGroupsDict(new_group, alreadyExtended=0)
            parser.saveLimitsGroupsDictYAML(parser.calInputLimitsGroupsDict, yaml_path)

    return {
        "status": "saved",
        "group_name": new_group.name,
        "limits_count": len(new_group.limitsRowMatchDictList),
        "yaml_path": yaml_path,
    }


@router.get("/series/{limit_name}")
def get_statistical_series(limit_name: str, limits_type: str = "burn_in"):
    """Get histogram data for a specific limit (statistical series).

    Returns the raw values and computed bins for visualization.
    Equivalent to the 'Statistical Series' button in Qt GUI.
    """
    if _last_controller is None:
        raise HTTPException(status_code=422, detail="No limits computed. Call POST /compute first.")

    # Find the row
    rows = _last_controller.limitsRowListDict.get("limitsRowList", [])
    target_row = None
    for row in rows:
        if row.limitName == limit_name and row.protType == ("BurnIn" if limits_type == "burn_in" else "Calibration"):
            target_row = row
            break

    if target_row is None:
        raise HTTPException(status_code=404, detail=f"Limit '{limit_name}' not found.")

    # Collect plot data (same logic as _collect_plot_data in stats_handlers.py)
    parsed_dict = (
        _last_controller.burnInParsedProtDict
        if limits_type == "burn_in"
        else _last_controller.calParsedProtDict
    )
    parsed_list = parsed_dict.get("parsedDataList", [])
    
    values = []
    protocol_indices = []
    for idx, parsed_data in enumerate(parsed_list):
        for diag_idx, diagram_dict in enumerate(parsed_data.diagrams):
            for param_name, param_obj in diagram_dict.items():
                if hasattr(param_obj, limit_name):
                    val = getattr(param_obj, limit_name)
                    if val is not None:
                        values.append(float(val))
                        protocol_indices.append(idx)

    if not values:
        raise HTTPException(status_code=404, detail=f"No data points found for limit '{limit_name}'.")

    # Compute histogram bins (using numpy-like logic)
    import math
    n = len(values)
    bin_count = min(int(math.sqrt(n)), 50)  # Sqrt rule with cap at 50 bins
    min_val = min(values)
    max_val = max(values)
    bin_width = (max_val - min_val) / bin_count if bin_count > 0 else 1.0
    
    bins = [min_val + i * bin_width for i in range(bin_count + 1)]
    hist = [0] * bin_count
    
    for v in values:
        bin_idx = min(int((v - min_val) / bin_width), bin_count - 1)
        hist[bin_idx] += 1

    return {
        "limit_name": limit_name,
        "limits_type": limits_type,
        "values": values,
        "protocol_indices": protocol_indices,
        "count": len(values),
        "histogram": {
            "bins": bins,
            "counts": hist,
        },
        "statistics": {
            "mean": target_row.mean,
            "min": target_row.min,
            "max": target_row.max,
            "stDev": target_row.stDev,
            "lowerLimit": target_row.lowerLimit,
            "upperLimit": target_row.upperLimit,
        },
    }


class ExportPptxRequest(BaseModel):
    """Request to export current limits table to PPTX."""
    limits_type: str = "burn_in"
    output_filename: str = "statistics_export.pptx"


@router.post("/export/pptx")
def export_limits_pptx(request: ExportPptxRequest):
    """Export limits table to PPTX file.

    Creates a PowerPoint with statistics table (similar to group PPTX export in Qt GUI).
    """
    if _last_controller is None:
        raise HTTPException(status_code=422, detail="No limits computed. Call POST /compute first.")

    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        import shared.auxiliary as aux
        from datetime import datetime

        # Get rows
        rows = _last_controller.limitsRowListDict.get("limitsRowList", [])
        prot_type = "BurnIn" if request.limits_type == "burn_in" else "Calibration"
        filtered_rows = [r for r in rows if r.protType == prot_type]

        if not filtered_rows:
            raise HTTPException(status_code=404, detail=f"No {prot_type} limits found.")

        # Create presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = f"{prot_type} Limits Statistics — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        title_frame.paragraphs[0].font.size = Pt(18)
        title_frame.paragraphs[0].font.bold = True

        # Add table
        rows_count = len(filtered_rows) + 1  # +1 for header
        cols_count = 10  # limitName, protType, lower, upper, mean, min, max, stDev, maxDev, cg
        
        table_shape = slide.shapes.add_table(
            rows_count, cols_count,
            Inches(0.3), Inches(1.0),
            Inches(9.4), Inches(5.5)
        )
        table = table_shape.table

        # Header row
        headers = ["Limit Name", "Type", "Lower", "Upper", "Mean", "Min", "Max", "StDev", "MaxDev", "Cg"]
        for col_idx, header in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = header
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(9)

        # Data rows
        for row_idx, row_obj in enumerate(filtered_rows, start=1):
            values = [
                row_obj.limitName,
                row_obj.protType,
                f"{row_obj.lowerLimit:.2f}",
                f"{row_obj.upperLimit:.2f}",
                f"{row_obj.mean:.3f}",
                f"{row_obj.min:.3f}",
                f"{row_obj.max:.3f}",
                f"{row_obj.stDev:.3f}",
                f"{row_obj.maxDev:.3f}",
                f"{row_obj.cg:.2f}",
            ]
            for col_idx, val in enumerate(values):
                cell = table.cell(row_idx, col_idx)
                cell.text = str(val)
                cell.text_frame.paragraphs[0].font.size = Pt(8)

        # Save
        output_dir = aux.statisticOutputPath
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, request.output_filename)
        prs.save(output_path)

        return {
            "status": "exported",
            "output_path": output_path,
            "row_count": len(filtered_rows),
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to export PPTX: {str(e)}")